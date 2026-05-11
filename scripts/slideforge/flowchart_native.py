"""Render a small subset of Mermaid flowcharts as native PowerPoint shapes.

When the diagram matches this subset, users can edit box text, move shapes,
and adjust connectors in PowerPoint. Unsupported syntax falls back to PNG
rendering in ``pptx_builder``.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Set, Tuple

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# IR
# ---------------------------------------------------------------------------


@dataclass
class FlowGraph:
    direction: str  # LR TB RL BT
    nodes: Dict[str, str] = field(default_factory=dict)  # id -> label
    edges: List[Tuple[str, str]] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------

_UNSUPPORTED = re.compile(
    r"(?im)^(subgraph|end|classDiagram|sequenceDiagram|stateDiagram|erDiagram|"
    r"gantt|pie|journey|mindmap|timeline|gitGraph|C4Context|sankey|flowchart\s+TD)\b"
)

_HEADER_RE = re.compile(r"^\s*(flowchart|graph)\s+(LR|RL|TB|BT)\s*$", re.I)

_ARROWS = ("-.->", "-->", "===", "==>", "---", "--")

# Node-only line: id[label] or id(label) or id{label} etc.
_NODE_ONLY_RE = re.compile(
    r"^\s*(\w+)\s*"
    r"(\[[^\]\n]+\]|\([^)\n]+\)|\{[^}\n]+\}|\(\([^)\n]+\)\))\s*$"
)


def parse_simple_flowchart(code: str) -> Optional[FlowGraph]:
    """Parse ``code`` into a :class:`FlowGraph`, or ``None`` if unsupported."""
    raw = code.strip()
    if not raw or _UNSUPPORTED.search(raw):
        return None

    lines: List[str] = []
    for line in raw.splitlines():
        s = line.strip()
        if not s or s.startswith("%%"):
            continue
        if ";" in s and not s.startswith("%%"):
            for part in s.split(";"):
                p = part.strip()
                if p:
                    lines.append(p)
        else:
            lines.append(s)

    if not lines:
        return None

    hdr = _HEADER_RE.match(lines[0])
    if not hdr:
        return None
    direction = hdr.group(2).upper()
    body = lines[1:]

    nodes: Dict[str, str] = {}
    edges: List[Tuple[str, str]] = []

    for line in body:
        if _HEADER_RE.match(line):
            continue
        if re.match(r"(?i)^\s*(linkStyle|style|classDef|class)\b", line):
            continue

        m_edge = _split_edge_line(line)
        if m_edge is not None:
            left_s, right_s = m_edge
            pa = _parse_node_token(left_s)
            pb = _parse_node_token(right_s)
            if pa is None or pb is None:
                return None
            a, la = pa
            b, lb = pb
            if a == b:
                return None
            edges.append((a, b))
            _merge_node_label(nodes, a, la)
            _merge_node_label(nodes, b, lb)
            continue

        m_node = _NODE_ONLY_RE.match(line)
        if m_node:
            nid, wrapped = m_node.group(1), m_node.group(2)
            label = _unwrap_label(wrapped)
            if not label:
                return None
            nodes[nid] = label
            continue
        return None

    if not edges:
        return None

    for a, b in edges:
        nodes.setdefault(a, a)
        nodes.setdefault(b, b)

    if len(nodes) > 24:
        return None

    return FlowGraph(direction=direction, nodes=nodes, edges=edges)


def _merge_node_label(nodes: Dict[str, str], nid: str, label: str) -> None:
    """Prefer explicit bracket labels; do not replace them with a bare id token."""
    if nid not in nodes or label != nid:
        nodes[nid] = label


def _unwrap_label(wrapped: str) -> str:
    inner = wrapped[1:-1]
    inner = inner.strip()
    if inner.startswith("(") and inner.endswith(")"):
        inner = inner[1:-1].strip()
    return inner.strip()


def _split_edge_line(line: str) -> Optional[Tuple[str, str]]:
    """Split ``line`` on first supported arrow, stripping ``|edge label|``."""
    s = line.strip()
    m_pipe = re.search(r"\|[^|]+\|", s)
    if m_pipe:
        s = (s[: m_pipe.start()] + s[m_pipe.end() :]).replace("  ", " ").strip()
    for arrow in _ARROWS:
        if arrow in s:
            i = s.index(arrow)
            left, right = s[:i].strip(), s[i + len(arrow) :].strip()
            if left and right:
                return left, right
    return None


def _parse_node_token(s: str) -> Optional[Tuple[str, str]]:
    """Return ``(id, label)`` for a mermaid node token."""
    s = s.strip()
    if not s:
        return None
    m = re.match(r"^(\w+)$", s)
    if m:
        w = m.group(1)
        return w, w
    m = re.match(r"^(\w+)\[([^\]]*)\]\s*$", s)
    if m:
        w, lab = m.group(1), m.group(2).strip()
        return w, lab or w
    m = re.match(r"^(\w+)\(([^)]*)\)\s*$", s)
    if m:
        w, lab = m.group(1), m.group(2).strip()
        return w, lab or w
    m = re.match(r"^(\w+)\{([^}]*)\}\s*$", s)
    if m:
        w, lab = m.group(1), m.group(2).strip()
        return w, lab or w
    m = re.match(r"^(\w+)\(\(([^)]*)\)\)\s*$", s)
    if m:
        w, lab = m.group(1), m.group(2).strip()
        return w, lab or w
    return None


# ---------------------------------------------------------------------------
# Layout + draw
# ---------------------------------------------------------------------------


def _layer_ranks(edges: List[Tuple[str, str]], nodes: Set[str]) -> Dict[str, int]:
    rank: Dict[str, int] = {n: 0 for n in nodes}
    for _ in range(len(nodes) + 2):
        changed = False
        for u, v in edges:
            if rank[v] < rank[u] + 1:
                rank[v] = rank[u] + 1
                changed = True
        if not changed:
            break
    return rank


def _group_layers(rank: Dict[str, int]) -> Dict[int, List[str]]:
    layers: Dict[int, List[str]] = {}
    for nid, r in rank.items():
        layers.setdefault(r, []).append(nid)
    for r in layers:
        layers[r].sort()
    return layers


def try_draw_mermaid_flowchart(
    slide,
    code: str,
    left: int,
    top: int,
    width: int,
    height: int,
    *,
    font_name: Optional[str] = None,
) -> bool:
    """If ``code`` is a supported flowchart, draw native shapes and return True."""
    g = parse_simple_flowchart(code)
    if g is None:
        return False

    nodes_set = set(g.nodes.keys())
    rank = _layer_ranks(g.edges, nodes_set)
    layers = _group_layers(rank)
    layer_ids = sorted(layers.keys())
    if not layer_ids:
        return False

    pad_x = int(width * 0.04)
    pad_y = int(height * 0.06)
    inner_w = max(1, width - 2 * pad_x)
    inner_h = max(1, height - 2 * pad_y)
    n_layers = len(layer_ids)
    max_stack = max(len(layers[r]) for r in layer_ids)

    if g.direction in ("LR", "RL"):
        layer_pitch = inner_w / max(1, n_layers)
        stack_pitch = inner_h / max(1, max_stack)
        node_w = int(min(layer_pitch * 0.72, inner_w / max(2, n_layers)))
        node_h = int(min(stack_pitch * 0.78, inner_h / max(2, max_stack)))
    else:
        layer_pitch = inner_h / max(1, n_layers)
        stack_pitch = inner_w / max(1, max_stack)
        node_w = int(min(stack_pitch * 0.78, inner_w / max(2, max_stack)))
        node_h = int(min(layer_pitch * 0.72, inner_h / max(2, n_layers)))

    node_w = max(int(Emu(Inches(0.55))), node_w)
    node_h = max(int(Emu(Inches(0.28))), node_h)

    pos: Dict[str, Tuple[int, int, int, int]] = {}

    for li, r in enumerate(layer_ids):
        stack = layers[r]
        m = len(stack)
        for si, nid in enumerate(stack):
            if g.direction == "LR":
                cx = left + pad_x + (li + 0.5) * (inner_w / max(1, n_layers))
                cy = top + pad_y + (si + 0.5) * (inner_h / max(1, m))
                x1 = int(cx - node_w / 2)
                y1 = int(cy - node_h / 2)
            elif g.direction == "RL":
                cx = left + pad_x + (n_layers - 1 - li + 0.5) * (inner_w / max(1, n_layers))
                cy = top + pad_y + (si + 0.5) * (inner_h / max(1, m))
                x1 = int(cx - node_w / 2)
                y1 = int(cy - node_h / 2)
            elif g.direction == "TB":
                cx = left + pad_x + (si + 0.5) * (inner_w / max(1, m))
                cy = top + pad_y + (li + 0.5) * (inner_h / max(1, n_layers))
                x1 = int(cx - node_w / 2)
                y1 = int(cy - node_h / 2)
            else:  # BT
                cx = left + pad_x + (si + 0.5) * (inner_w / max(1, m))
                cy = top + pad_y + (n_layers - 1 - li + 0.5) * (inner_h / max(1, n_layers))
                x1 = int(cx - node_w / 2)
                y1 = int(cy - node_h / 2)
            pos[nid] = (x1, y1, node_w, node_h)

    # White backing so placeholders / grid do not show through.
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg.line.fill.background()

    for u, v in g.edges:
        if u not in pos or v not in pos:
            continue
        x1, y1, w1, h1 = pos[u]
        x2, y2, w2, h2 = pos[v]
        ax, ay = _anchor_out(u, v, pos, g.direction)
        bx, by = _anchor_in(u, v, pos, g.direction)
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ax, ay, bx, by)
        conn.line.color.rgb = RGBColor(0x55, 0x55, 0x55)
        conn.line.width = Pt(1.0)

    for nid, (x1, y1, nw, nh) in pos.items():
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x1, y1, nw, nh)
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(0xF2, 0xF6, 0xFC)
        sh.line.color.rgb = RGBColor(0x2F, 0x55, 0x99)
        sh.line.width = Pt(1)
        tf = sh.text_frame
        tf.clear()
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(50000)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = g.nodes.get(nid, nid)
        run.font.size = Pt(10 if nw < Emu(Inches(1.2)) else 11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        if font_name:
            run.font.name = font_name

    return True


def _anchor_out(
    u: str, v: str, pos: Dict[str, Tuple[int, int, int, int]], direction: str
) -> Tuple[int, int]:
    x, y, w, h = pos[u]
    cx, cy = x + w // 2, y + h // 2
    if direction == "LR":
        return x + w, cy
    if direction == "RL":
        return x, cy
    if direction == "TB":
        return cx, y + h
    return cx, y


def _anchor_in(
    u: str, v: str, pos: Dict[str, Tuple[int, int, int, int]], direction: str
) -> Tuple[int, int]:
    x, y, w, h = pos[v]
    cx, cy = x + w // 2, y + h // 2
    if direction == "LR":
        return x, cy
    if direction == "RL":
        return x + w, cy
    if direction == "TB":
        return cx, y
    return cx, y + h
